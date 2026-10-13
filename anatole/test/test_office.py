"""Word export: the document, its equations and units, and the problems it reports."""

import importlib.util
import zipfile
from pathlib import Path

import pytest

from anatole import Notation, NotationSet, Rendered, Translated
from anatole.cli import main
from anatole.export import export, to_docx
from anatole.export.office import is_available

pytestmark = pytest.mark.skipif(not is_available(), reason="needs pandoc (anatole[office])")

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "notations.toml"
has_docx = importlib.util.find_spec("docx") is not None


def document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").decode()


@pytest.mark.skipif(not has_docx, reason="needs python-docx")
def test_lists(tmp_path: Path, example: NotationSet) -> None:
    import docx

    path = tmp_path / "notations.docx"
    assert to_docx(example, path, language="fr") == []
    document = docx.Document(str(path))
    headings = [(p.style.name if p.style else "", p.text) for p in document.paragraphs]
    assert headings == [
        ("Heading 1", "Nomenclature"),
        ("Heading 2", "Symboles"),
        ("Heading 3", "Lettres latines"),
        ("Heading 3", "Lettres grecques"),
        ("Heading 3", "Indices"),
        ("Heading 2", "Opérateurs"),
        ("Heading 1", "Acronyms"),
        ("Heading 2", "Acronymes"),
        ("Heading 2", "Abréviations"),
    ]
    rows = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
    assert ["", "Vitesse", "m\xa0s\u22121"] in rows  # siunitx, converted
    assert ["", "Masse volumique", "kg\xa0m\u22123"] in rows
    assert ["MFN", "mécanique des fluides numérique"] in rows
    assert ["p.r.à", "par rapport à"] in rows


def test_symbols_are_equations(tmp_path: Path, example: NotationSet) -> None:
    path = tmp_path / "notations.docx"
    _ = to_docx(example, path)
    xml = document_xml(path)
    # 8 symbols, and delta in the description of redelta.
    assert xml.count("<m:oMath>") == 9
    assert '<m:sty m:val="bi" /></m:rPr><m:t>u</m:t>' in xml  # the vector style: bold
    assert "<m:t>δ</m:t>" in xml


def test_problems_by_entry(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    notations = NotationSet(
        [
            Notation(
                "length",
                Rendered.both("L"),
                Translated({"en": "Length"}),
                Rendered.both(r"\si{\furlong}"),
            ),
            Notation("weird", Rendered(r"\unknowncmd{x}", "x"), Translated({"en": "Weird"})),
        ]
    )
    with caplog.at_level("WARNING", logger="anatole"):
        problems = to_docx(notations, tmp_path / "bad.docx")
    assert problems == [
        r"notation 'length': Word: \furlong is not converted (left out)",
        r"notation 'weird': Word: the math \unknowncmd{x} is not converted (kept as TeX)",
    ]
    assert caplog.messages == problems
    assert (tmp_path / "bad.docx").exists()  # written anyway


@pytest.mark.skipif(not has_docx, reason="needs python-docx")
def test_acronym_reference_prints_the_short_form(tmp_path: Path, example: NotationSet) -> None:
    import docx

    example.add(Notation("solver", Rendered.both("S"), Translated({"en": "The @{cfd} solver"})))
    path = tmp_path / "notations.docx"
    _ = to_docx(example, path)
    cells = [c.text for t in docx.Document(str(path)).tables for r in t.rows for c in r.cells]
    assert "The CFD solver" in cells


def test_reference_doc_and_export(tmp_path: Path, example: NotationSet) -> None:
    reference = tmp_path / "reference.docx"
    _ = to_docx(example, reference)
    path = tmp_path / "with-styles.docx"
    assert export(example, path, reference_doc=reference, tags=["nothing"]) == []
    xml = document_xml(path)
    assert "large-eddy simulation" not in xml  # les is tagged turbulence
    assert "computational fluid dynamics" in xml


def test_command_line(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    path = tmp_path / "notations.docx"
    assert main(["export", "-l", "fr", str(EXAMPLE), str(path)]) == 0
    assert capsys.readouterr().out == f"{path}: written\n"
    assert "Vitesse" in document_xml(path)
